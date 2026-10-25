from __future__ import annotations

import logging
from datetime import date
from io import BytesIO
from typing import Any

from app.core.config import Environment, settings
from app.schemas.candidate import coerce_named_items
from app.services.gemini_service import (
    GeminiError,
    analyze_resume,
    extract_text_from_file,
    is_configured,
)
from app.services.resume_files import MIME_DOCX, MIME_JPEG, MIME_PDF, MIME_PNG

logger = logging.getLogger(__name__)

DOCX_MIN_CHARS = 40
SAFE_OCR_ERROR = "Could not read this resume. Try a clearer PDF or image."
SAFE_GEMINI_ERROR = "Could not analyze this resume. Try again in a moment."
SAFE_DOCX_ERROR = "This Word file has too little text. Upload a PDF, PNG, or JPG instead."


class ResumeParseError(RuntimeError):
    def __init__(self, message: str, *, stage: str) -> None:
        super().__init__(message)
        self.stage = stage


def empty_parse_result() -> dict[str, Any]:
    return {
        "given_name": "",
        "middle_name": "",
        "family_name": "",
        "email": "",
        "phone": "",
        "address": "",
        "city": "",
        "country": "",
        "current_role": "",
        "summary": "",
        "experience": [],
        "education": [],
        "skills": [],
        "technologies": [],
        "projects": [],
        "certifications": [],
        "achievements": [],
    }


def _use_document_ai() -> bool:
    return bool(settings.document_ai_processor_id) and settings.environment is not Environment.LOCAL


def extract_text(data: bytes, content_type: str) -> str:
    if content_type == MIME_DOCX:
        return _extract_docx(data)
    if content_type in {MIME_PDF, MIME_PNG, MIME_JPEG}:
        if _use_document_ai():
            try:
                return _extract_document_ai(data, content_type)
            except ResumeParseError:
                logger.info("resume_document_ai_fallback content_type=%s", content_type)
        if content_type == MIME_PDF:
            text = _extract_pdf_local(data, required=False)
            if text:
                return text
        return _extract_gemini_ocr(data, content_type)
    raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr")


def analyze_resume_text(text: str) -> dict[str, Any]:
    if not text.strip():
        return empty_parse_result()
    if not is_configured():
        return empty_parse_result()
    try:
        return normalize_parse_result(analyze_resume(text).model_dump())
    except Exception:
        logger.exception("resume_gemini_failed")
        raise ResumeParseError(SAFE_GEMINI_ERROR, stage="gemini") from None


def parse_resume_bytes(data: bytes, content_type: str) -> tuple[str, dict[str, Any]]:
    text = extract_text(data, content_type)
    parsed = analyze_resume_text(text)
    return text, parsed


def has_parse_evidence(parse: dict[str, Any] | None) -> bool:
    if not isinstance(parse, dict):
        return False
    if str(parse.get("summary") or "").strip():
        return True
    if str(parse.get("given_name") or "").strip() or str(parse.get("family_name") or "").strip():
        return True
    for key in ("skills", "experience", "projects"):
        value = parse.get(key)
        if isinstance(value, list) and value:
            return True
    return False


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text).strip()
    if len(text) < DOCX_MIN_CHARS:
        raise ResumeParseError(SAFE_DOCX_ERROR, stage="ocr")
    return text


def _extract_pdf_local(data: bytes, *, required: bool = True) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except ResumeParseError:
        raise
    except Exception:
        logger.exception("resume_pdf_local_failed")
        if required:
            raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr") from None
        return ""
    if text:
        return text
    if required:
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr")
    return ""


def _extract_gemini_ocr(data: bytes, mime_type: str) -> str:
    if not is_configured():
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr")
    try:
        text = extract_text_from_file(data, mime_type).strip()
    except ResumeParseError:
        raise
    except (GeminiError, Exception):
        logger.exception("resume_gemini_ocr_failed")
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr") from None
    if not text:
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr")
    return text


def _extract_document_ai(data: bytes, mime_type: str) -> str:
    try:
        from google.cloud import documentai

        client = documentai.DocumentProcessorServiceClient()
        name = client.processor_path(
            settings.google_cloud_project,
            settings.vertex_location,
            settings.document_ai_processor_id,
        )
        request = documentai.ProcessRequest(
            name=name,
            raw_document=documentai.RawDocument(content=data, mime_type=mime_type),
        )
        result = client.process_document(request=request)
        text = (result.document.text or "").strip()
    except ResumeParseError:
        raise
    except Exception:
        logger.exception("resume_ocr_failed")
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr") from None
    if not text:
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr")
    return text


def normalize_parse_result(payload: dict[str, Any] | None) -> dict[str, Any]:
    base = empty_parse_result()
    if not payload:
        return base
    for key in (
        "given_name",
        "middle_name",
        "family_name",
        "email",
        "phone",
        "address",
        "city",
        "country",
        "current_role",
        "summary",
    ):
        value = payload.get(key)
        base[key] = str(value).strip() if value is not None else ""
    base["experience"] = _list_of_dicts(payload.get("experience"))
    base["education"] = _list_of_dicts(payload.get("education"))
    base["skills"] = coerce_named_items(payload.get("skills"), "skill")
    base["technologies"] = coerce_named_items(payload.get("technologies"), "name")
    base["projects"] = _list_of_dicts(payload.get("projects"))
    base["certifications"] = _list_of_dicts(payload.get("certifications"))
    base["achievements"] = _list_of_dicts(payload.get("achievements"))
    return base


def resume_public_view(resume: dict) -> dict:
    processing = resume.get("processing") if isinstance(resume.get("processing"), dict) else {}
    parse_result = resume.get("parse_result") if isinstance(resume.get("parse_result"), dict) else None
    return {
        "text": resume.get("text") or "",
        "url": resume.get("url") or "",
        "filename": resume.get("filename") or "",
        "content_type": resume.get("content_type") or "",
        "size_bytes": int(resume.get("size_bytes") or 0),
        "processing": {
            "status": processing.get("status") or "",
            "error": processing.get("error") or "",
            "started_at": processing.get("started_at") or "",
            "finished_at": processing.get("finished_at") or "",
            "applied_at": processing.get("applied_at") or "",
        },
        "parse_result": normalize_parse_result(parse_result) if parse_result else None,
    }


def parse_optional_date(value: object) -> date | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _list_of_dicts(value: object) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]
