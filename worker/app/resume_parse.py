from __future__ import annotations

import json
import logging
import re
from io import BytesIO
from typing import Any

from pydantic import BaseModel, ConfigDict, Field, ValidationError

from app.config import Environment, settings

logger = logging.getLogger(__name__)

MIME_PDF = "application/pdf"
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_PNG = "image/png"
MIME_JPEG = "image/jpeg"
DOCX_MIN_CHARS = 40
SAFE_OCR_ERROR = "Could not read this resume. Try a clearer PDF or image."
SAFE_GEMINI_ERROR = "Could not analyze this resume. Try again in a moment."
SAFE_DOCX_ERROR = "This Word file has too little text. Upload a PDF, PNG, or JPG instead."
OCR_INSTRUCTION = (
    "Extract all readable text from this resume. Return plain text only. Do not summarize."
)

PARSE_INSTRUCTION = """Extract a candidate profile from the resume text.
Return JSON only with this exact shape:
{
  "given_name": "",
  "middle_name": "",
  "family_name": "",
  "email": "",
  "phone": "",
  "address": "",
  "city": "",
  "country": "",
  "current_role": "",
  "summary": "",
  "experience": [{"title": "", "company_name": "", "start_date": "", "end_date": "", "description": ""}],
  "education": [{"school": "", "degree": "", "field_of_study": "", "start_date": "", "end_date": ""}],
  "skills": [{"skill": "", "years": null}],
  "technologies": [{"name": "", "years": null}],
  "projects": [{"name": "", "description": "", "url": ""}],
  "certifications": [{"name": "", "issuer": "", "issued_on": ""}],
  "achievements": [{"title": "", "description": "", "occurred_on": ""}]
}
Rules:
- Use ISO dates (YYYY-MM-DD) when possible. Use null for unknown years. Empty string when unknown.
- Infer skill and technology years from employment date ranges when the resume does not state years.
- Include skills implied by experience and project descriptions, not only an explicit skills list.
- Skills are capabilities. Technologies are tools, languages, or platforms.
- Normalize common names (React.js → React, TS → TypeScript, Postgres → PostgreSQL, Node.js → Node.js).
- Split city and country when an address contains both.
- Do not invent employers, schools, or projects that are not in the text.
"""


class _GeminiModel(BaseModel):
    model_config = ConfigDict(extra="ignore")


class ResumeExperience(_GeminiModel):
    title: str = ""
    company_name: str = ""
    start_date: str = ""
    end_date: str = ""
    description: str = ""


class ResumeEducation(_GeminiModel):
    school: str = ""
    degree: str = ""
    field_of_study: str = ""
    start_date: str = ""
    end_date: str = ""


class ResumeSkill(_GeminiModel):
    skill: str = ""
    years: float | None = None


class ResumeTechnology(_GeminiModel):
    name: str = ""
    years: float | None = None


class ResumeProject(_GeminiModel):
    name: str = ""
    description: str = ""
    url: str = ""


class ResumeCertification(_GeminiModel):
    name: str = ""
    issuer: str = ""
    issued_on: str = ""


class ResumeAchievement(_GeminiModel):
    title: str = ""
    description: str = ""
    occurred_on: str = ""


class ResumeAnalysis(_GeminiModel):
    given_name: str = ""
    middle_name: str = ""
    family_name: str = ""
    email: str = ""
    phone: str = ""
    address: str = ""
    city: str = ""
    country: str = ""
    current_role: str = ""
    summary: str = ""
    experience: list[ResumeExperience] = Field(default_factory=list)
    education: list[ResumeEducation] = Field(default_factory=list)
    skills: list[ResumeSkill] = Field(default_factory=list)
    technologies: list[ResumeTechnology] = Field(default_factory=list)
    projects: list[ResumeProject] = Field(default_factory=list)
    certifications: list[ResumeCertification] = Field(default_factory=list)
    achievements: list[ResumeAchievement] = Field(default_factory=list)


class ResumeParseError(RuntimeError):
    def __init__(self, message: str, *, stage: str) -> None:
        super().__init__(message)
        self.stage = stage


def empty_parse_result() -> dict[str, Any]:
    return {
        "given_name": "",
        "middle_name": "",
        "family_name": "",
        "email": "",
        "phone": "",
        "address": "",
        "city": "",
        "country": "",
        "current_role": "",
        "summary": "",
        "experience": [],
        "education": [],
        "skills": [],
        "technologies": [],
        "projects": [],
        "certifications": [],
        "achievements": [],
    }


def _use_document_ai() -> bool:
    return bool(settings.document_ai_processor_id) and settings.environment is not Environment.LOCAL


def _gemini_api_key() -> str | None:
    if settings.gemini_api_key is None:
        return None
    return settings.gemini_api_key.get_secret_value()


def _build_gemini_client():
    from google import genai

    if settings.use_vertex_ai:
        return genai.Client(
            vertexai=True,
            project=settings.google_cloud_project,
            location=settings.vertex_location,
        )
    key = _gemini_api_key()
    if key:
        return genai.Client(api_key=key)
    raise ResumeParseError(SAFE_GEMINI_ERROR, stage="gemini")


def extract_text(data: bytes, content_type: str) -> str:
    if content_type == MIME_DOCX:
        from docx import Document

        document = Document(BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text).strip()
        if len(text) < DOCX_MIN_CHARS:
            raise ResumeParseError(SAFE_DOCX_ERROR, stage="ocr")
        return text
    if content_type in {MIME_PDF, MIME_PNG, MIME_JPEG}:
        if _use_document_ai():
            return _extract_document_ai(data, content_type)
        if content_type == MIME_PDF:
            text = _extract_pdf_local(data)
            if text:
                return text
        return _extract_gemini_ocr(data, content_type)
    raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr")


def analyze_resume_text(text: str) -> dict[str, Any]:
    if not text.strip() or not settings.gemini_configured:
        return empty_parse_result()
    try:
        from google.genai import types

        client = _build_gemini_client()
        response = client.models.generate_content(
            model=settings.gemini_model or "gemini-3.5-flash",
            contents=f"{PARSE_INSTRUCTION}\n{text[:20000]}",
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=ResumeAnalysis,
            ),
        )
        raw = (response.text or "").strip()
        try:
            parsed = ResumeAnalysis.model_validate_json(raw)
            return normalize_parse_result(parsed.model_dump())
        except ValidationError:
            return normalize_parse_result(_loads_json(raw))
    except ResumeParseError:
        raise
    except Exception:
        logger.exception("resume_gemini_failed")
        raise ResumeParseError(SAFE_GEMINI_ERROR, stage="gemini") from None


def parse_resume_bytes(data: bytes, content_type: str) -> tuple[str, dict[str, Any]]:
    text = extract_text(data, content_type)
    return text, analyze_resume_text(text)


def normalize_parse_result(payload: dict[str, Any] | None) -> dict[str, Any]:
    base = empty_parse_result()
    if not payload:
        return base
    for key in (
        "given_name",
        "middle_name",
        "family_name",
        "email",
        "phone",
        "address",
        "city",
        "country",
        "current_role",
        "summary",
    ):
        value = payload.get(key)
        base[key] = str(value).strip() if value is not None else ""
    for key in (
        "experience",
        "education",
        "projects",
        "certifications",
        "achievements",
    ):
        value = payload.get(key)
        base[key] = [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []
    base["skills"] = _named_items(payload.get("skills"), "skill")
    base["technologies"] = _named_items(payload.get("technologies"), "name")
    return base


def _named_items(value: object, name_key: str) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    items: list[dict[str, Any]] = []
    for item in value:
        if isinstance(item, str) and item.strip():
            items.append({name_key: item.strip()})
        elif isinstance(item, dict):
            items.append(item)
    return items


def _extract_pdf_local(data: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception:
        logger.exception("resume_pdf_local_failed")
        return ""


def _extract_gemini_ocr(data: bytes, mime_type: str) -> str:
    if not settings.gemini_configured:
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr")
    try:
        from google.genai import types

        client = _build_gemini_client()
        response = client.models.generate_content(
            model=settings.gemini_model or "gemini-3.5-flash",
            contents=[
                types.Part.from_bytes(data=data, mime_type=mime_type),
                OCR_INSTRUCTION,
            ],
        )
        text = (response.text or "").strip()
    except ResumeParseError:
        raise
    except Exception:
        logger.exception("resume_gemini_ocr_failed")
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr") from None
    if not text:
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr")
    return text


def _extract_document_ai(data: bytes, mime_type: str) -> str:
    from google.cloud import documentai

    try:
        client = documentai.DocumentProcessorServiceClient()
        name = client.processor_path(
            settings.google_cloud_project,
            settings.vertex_location,
            settings.document_ai_processor_id,
        )
        result = client.process_document(
            request=documentai.ProcessRequest(
                name=name,
                raw_document=documentai.RawDocument(content=data, mime_type=mime_type),
            )
        )
        text = (result.document.text or "").strip()
    except ResumeParseError:
        raise
    except Exception:
        logger.exception("resume_ocr_failed")
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr") from None
    if not text:
        raise ResumeParseError(SAFE_OCR_ERROR, stage="ocr")
    return text


def _loads_json(raw: str) -> dict[str, Any]:
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
        if not match:
            return empty_parse_result()
        payload = json.loads(match.group(0))
    return payload if isinstance(payload, dict) else empty_parse_result()
